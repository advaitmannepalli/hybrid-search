import requests
import pdfplumber
import io
from bs4 import BeautifulSoup
from opensearchpy import OpenSearch
from urllib.parse import urljoin, urlparse
from sentence_transformers import SentenceTransformer
import pytesseract
from pdf2image import convert_from_bytes
import threading
from queue import Queue, Empty
from concurrent.futures import ThreadPoolExecutor
from docx import Document
import openpyxl

client = OpenSearch(hosts=["http://localhost:9200"])
model = SentenceTransformer("all-MiniLM-L6-v2")

START_URL = "https://www.ercot.com"
MAX_PAGES = 500

def create_session():
    session = requests.Session()

    session.headers.update({
        "User-Agent": (
            "HybridSearchBot/1.0 "
            "(Educational Search Engine Project)"
        )
    })

    return session

def reset_index():
    if client.indices.exists(index="docs"):
        client.indices.delete(index="docs")
        print("Deleted old index")
    client.indices.create(index="docs", body={
        "settings": {
            "index.knn": True
        },
        "mappings": {
            "properties": {
                "title":     { "type": "text" },
                "body":      { "type": "text" },
                "url":       { "type": "keyword" },
                "embedding": {
                    "type": "knn_vector",
                    "dimension": 384
                }
            }
        }
    })
    print("Created fresh index with kNN enabled")

def get_links(url, soup):
    html_links = set()
    pdf_links = set()
    xlsx_links = set()
    docx_links = set()

    base_domain = urlparse(START_URL).netloc

    for a_tag in soup.find_all("a", href=True):
        href = a_tag["href"]
        full_url = urljoin(url, href)
        parsed = urlparse(full_url)

        if parsed.netloc != base_domain:
            continue

        clean_url = full_url.split("#")[0]

        if clean_url.endswith(".pdf"):
            pdf_links.add(clean_url)
        elif clean_url.endswith(".docx"):
            docx_links.add(clean_url)
        elif clean_url.endswith(".xlsx"):
            xlsx_links.add(clean_url)
        elif not any(clean_url.endswith(ext) for ext in [".zip", ".png", ".jpg", ".mp4"]):
            html_links.add(clean_url)

    return html_links, pdf_links, docx_links, xlsx_links

def fetch_page(url, session):
    response = session.get(url, timeout=10)
    soup = BeautifulSoup(response.text, "html.parser")
    for tag in soup(["nav", "footer", "script", "style"]):
        tag.decompose()
    title = soup.title.string if soup.title else url
    body = soup.get_text(separator=" ", strip=True)
    return title, body, soup

def chunk_text(text, size=300, overlap=30):
    words = text.split()
    chunks = []
    i = 0
    while i < len(words):
        chunk = " ".join(words[i:i+size])
        chunks.append(chunk)
        i += size - overlap
    return chunks

def index_chunks(title, chunks, url):
    for i, chunk in enumerate(chunks):
        if not chunk.strip():
            continue
        embedding = model.encode(chunk).tolist()
        doc = {
            "title":     f"{title} (part {i+1})",
            "body":      chunk,
            "url":       url,
            "embedding": embedding
        }
        client.index(index="docs", body=doc)

def ingest_url(url, session):
    title, body, soup = fetch_page(url, session)
    chunks = chunk_text(body)
    index_chunks(title, chunks, url)
    return soup

def ingest_pdf(url, session):
    print(f"  → reading PDF: {url}")
    response = session.get(url, timeout=10)

    # load the PDF from memory instead of saving to disk
    pdf_file = io.BytesIO(response.content)

    with pdfplumber.open(pdf_file) as pdf:
        title = url.split("/")[-1].replace(".pdf", "")
        full_text = ""
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                full_text += text + " "

    # if pdfplumber found nothing, fall back to OCR
    if not full_text.strip():
        print(f"  → no text found, trying OCR...")
        images = convert_from_bytes(response.content)
        for i, image in enumerate(images):
            text = pytesseract.image_to_string(image)
            if text:
                full_text += text + " "
            print(f"  → OCR processed page {i+1}/{len(images)}")

    if not full_text.strip():
        print(f"  → no text found even after OCR, skipping")
        return

    chunks = chunk_text(full_text)
    index_chunks(title, chunks, url)
    print(f"  → indexed {len(chunks)} chunks from PDF")

def ingest_docx(url, session):
    print(f"  → reading DOCX: {url}")
    response = session.get(url, timeout=10)

    doc = Document(io.BytesIO(response.content))
    title = url.split("/")[-1].replace(".docx", "")
    full_text = " ".join([para.text for para in doc.paragraphs if para.text.strip()])

    if not full_text.strip():
        print(f"  → no text found, skipping")
        return

    chunks = chunk_text(full_text)
    index_chunks(title, chunks, url)
    print(f"  → indexed {len(chunks)} chunks from DOCX")

def ingest_xlsx(url, session):
    print(f"  → reading XLSX: {url}")
    response = session.get(url, timeout=10)

    workbook = openpyxl.load_workbook(io.BytesIO(response.content), data_only=True)
    title = url.split("/")[-1].replace(".xlsx", "")
    full_text = ""

    for sheet in workbook.worksheets:
        # get headers from first row
        headers = []
        for cell in sheet[1]:
            headers.append(str(cell.value) if cell.value is not None else "")

        # convert each subsequent row to labeled text
        for row in sheet.iter_rows(min_row=2, values_only=True):
            # skip completely empty rows
            if not any(cell is not None for cell in row):
                continue

            row_text = " | ".join([
                f"{headers[i]}: {str(val)}"
                for i, val in enumerate(row)
                if val is not None and i < len(headers) and headers[i]
            ])

            if row_text.strip():
                full_text += row_text + "\n"

    if not full_text.strip():
        print(f"  → no text found, skipping")
        return

    chunks = chunk_text(full_text)
    index_chunks(title, chunks, url)
    print(f"  → indexed {len(chunks)} chunks from XLSX")

def crawl():
    url_queue = Queue() # no need for manual lock, Queues already have
    url_queue.put(START_URL)
    
    visited = set()
    visited_lock = threading.Lock()  # prevents two threads touching visited at same time
    
    def process_url():
        session = create_session()

        while True:
            try:
                # grab next URL from queue, give up after 3 seconds if empty
                url = url_queue.get(timeout=3)
            except Empty:
                break

            # thread safe check — skip if already visited or over limit
            with visited_lock:
                if url in visited or len(visited) >= MAX_PAGES:
                    url_queue.task_done()
                    continue
                visited.add(url)

            try:
                print(f"[{len(visited)}/{MAX_PAGES}] Crawling {url}")

                if url.endswith(".pdf"):
                    ingest_pdf(url, session)
                elif url.endswith(".docx"):
                    ingest_docx(url, session)
                elif url.endswith(".xlsx"):
                    ingest_xlsx(url, session)
                else:
                    soup = ingest_url(url, session)
                    new_links, new_pdfs, new_docx, new_xlsx = get_links(url, soup)

                    for link in new_links:
                        with visited_lock:
                            if link not in visited:
                                url_queue.put(link)

                    for pdf in new_pdfs:
                        with visited_lock:
                            if pdf not in visited:
                                url_queue.put(pdf)

                    for docx in new_docx:
                        with visited_lock:
                            if docx not in visited:
                                url_queue.put(docx)

                    for xlsx in new_xlsx:
                        with visited_lock:
                            if xlsx not in visited:
                                url_queue.put(xlsx)

            except Exception as e:
                print(f"  → skipping {url}: {e}")

            url_queue.task_done()

    # spin up 10 workers all running process_url simultaneously
    with ThreadPoolExecutor(max_workers=15) as executor:
        futures = [executor.submit(process_url) for _ in range(10)]

    print(f"\nDone! Crawled {len(visited)} pages.")

if __name__ == "__main__":
    reset_index()
    crawl()